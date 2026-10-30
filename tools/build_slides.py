"""生成题目三的汇报PPT（精炼版，约14页，5分钟讲得完）。

不重跑模型，只引用仓库里已有的图片与结果文件。
用法：python tools/build_slides.py
"""
from pathlib import Path
import struct

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables' / '题目三_汇报.pptx'
STEP02 = ROOT / 'figures' / 'step02'
STEP08 = ROOT / 'figures' / 'step08'
STEP09 = ROOT / 'figures' / 'step09'
FONT = 'Arial Unicode MS'

TITLE_BOX = (Inches(0.62), Inches(0.42), Inches(12.08), Inches(0.85))
BODY_BOX = (Inches(0.68), Inches(1.50), Inches(11.77), Inches(5.20))
PIC_BOX = (Inches(0.62), Inches(1.35), Inches(6.70), Inches(4.55))
NOTE_BOX = (Inches(7.55), Inches(1.45), Inches(4.90), Inches(4.90))
CAPTION_BOX = (Inches(0.68), Inches(6.05), Inches(11.77), Inches(0.75))
FOOTER_BOX = (Inches(0.62), Inches(6.95), Inches(10.90), Inches(0.30))
PAGE_BOX = (Inches(12.19), Inches(6.95), Inches(0.78), Inches(0.30))
FOOTER_TEXT = '题目三  无人设备电池健康评估与剩余寿命预测   2026.10.31'

PAGES = [
    {'title': '可视化：容量退化与个体差异', 'figure': STEP02 / '01_capacity_all_cycles.png',
     'notes': ['四块电池的容量—循环曲线，全部观测值。',
               'B0006 下降最快，B0007 最慢。',
               'B0018 只观测到 132 次，',
               '线条自然结束，不补齐、不插值。'],
     'caption': '数据：NASA 四块锂电池放电循环记录，共 636 行；纵轴为实测容量（Ah）。'},

    {'title': '统一窗口后的差异量化', 'figure': STEP02 / '03_loss_comparison.png',
     'notes': ['共同窗口取所有电池都有的前 132 次，',
               '保证比较公平。',
               '同样经历 131 个循环间隔：',
               'B0006 容量损失 35.36%，最多；',
               'B0007 损失 21.90%，最少。',
               '最多与最少相差 13.5 个百分点。'],
     'caption': '同一横坐标下的容量损失与平均每循环损失量，差异主要来自电池个体本身。'},

    {'title': '预测误差对比', 'figure': STEP08 / 'rmse_comparison.png',
     'notes': ['平均 RMSE：',
               '基线 0.014425 Ah',
               '回归 0.017370 Ah',
               '森林 0.058766 Ah。',
               '平均 MAE 排序相同。',
               '两个主模型都没有赢过',
               '不需要训练的朴素基线。'],
     'caption': '各项指标均由逐点预测独立重算；129 个测试目标，24 组指标核对一致。'},

    {'title': '森林为什么会输', 'figure': STEP08 / 'B0006_predictions.png',
     'notes': ['树的预测值被限制在训练标签范围内，',
               '无法外推到更低的容量水平。',
               '实测跌破训练范围后，',
               '预测仍停在 1.33 Ah 附近。',
               '改进方向：改预测差分 Q(t)−Q(t−1)，',
               '等价于「基线＋修正项」。',
               '本次结果如实保留，未回调参。'],
     'caption': 'B0006 的实测、基线与两种模型预测；纵轴为容量（Ah）。'},

    {'title': '不确定性一：误差方向与高估风险', 'figure': STEP09 / '01_error_risk.png',
     'notes': ['平均偏差（预测−实测）：',
               '回归 −0.0111 Ah，偏低估；',
               '森林 +0.0444 Ah，高估；',
               '基线 +0.0023 Ah。',
               '森林在三块电池的全部',
               '测试点上高估容量，',
               'B0006 最大高估 0.173 Ah。'],
     'caption': '高估容量比低估更危险，不能因为回归整体偏低估就断言它更安全。'},

    {'title': '不确定性二：提前预警与寿命外推', 'figure': STEP09 / '03_lifetime_sensitivity.png',
     'notes': ['首次跌破 80% 研究阈值：',
               '第 101、61、124、75 次循环，',
               '均早于测试段起点 135 / 106，',
               '单步预测无法验证提前预警。',
               '外推随历史长度剧烈变化：',
               'B0005 前 30% 估第 228.6 次、',
               '前 50% 估 115.1 次，实测 101 次。'],
     'caption': '该范围反映外推敏感性，不是置信区间；241 个滚动起点仅有 4 个真实事件。'},
]

TEXT_PAGES = [
    {'kind': 'cover',
     'title': '题目三  无人设备电池健康评估与剩余寿命预测',
     'subtitle': '基于 NASA 四块锂电池的容量退化分析与预测',
     'lines': ['组号：待填        组长：待填',
               '组员：待填        汇报日期：2026.10.31']},

    {'kind': 'text', 'title': '汇报内容', 'size': 19,
     'lines': ['四项基本任务的完成情况：',
               '① 退化曲线可视化与个体差异识别',
               '② 特征提取、两种模型构建与比较',
               '③ 提前预警与寿命外推场景下的不确定性与风险',
               '④ 使用、维护与更换策略的建议',
               '最后给出成员分工、复现方式与仓库链接。']},

    {'kind': 'text', 'title': '数据与清洗', 'size': 16,
     'lines': ['· 数据：NASA 四块锂电池放电循环记录，共 636 行，字段为编号、循环次数、容量（Ah）、温度（℃）。',
              '· 规模：B0005、B0006、B0007 各 168 次循环，B0018 仅 132 次；温度恒为 24℃。',
              '· 参考口径：数据不含厂家额定容量，一律以「首循环容量」为参考，不称为真实 SOH。',
              '· 清洗检查：列名与顺序、缺失 0、完全重复 0、同电池同循环冲突 0、容量为正、循环从 1 连续。',
              '· 缺失一律不填补；填 0 等于伪造观测值。',
              '· 保留局部回升（B0005 22 次、B0006 18 次、B0007 34 次、B0018 21 次），不据此解释原因。',
              '· 统计跳变：相邻变化量的 IQR 规则标记待核查 58 条、极端 30 条，阈值只用前 60% 数据估计；',
              '  限幅版仅作对照，模型继续使用原值保留版。']},

    {'kind': 'text', 'title': '特征与时间划分', 'size': 16,
     'lines': ['· 任务：预测第 t 次循环容量时，只允许使用第 t 次之前已经测到的容量。',
              '· 特征：上一次容量 lag1、前 5 次均值 mean5、前 5 次最小二乘斜率 slope5、循环序号 cycle。',
              '· 防泄漏：先 shift(1) 再 rolling(5)，滚动窗口绝不越过当前预测目标。',
              '· 时间划分：每块电池按时间顺序 60% 训练、20% 验证、20% 测试，不打乱顺序。',
              '· 例：B0005 训练第 6–100 次、验证 101–134 次、测试 135–168 次。',
              '· 丢弃每块电池前 5 条没有完整历史窗口的记录，共 616 条样本、129 个测试目标。',
              '· 测试前在训练加验证段重新拟合，测试估计器固定不动。']},

    {'kind': 'text', 'title': '三种方法', 'size': 16,
     'lines': ['· 容量保持基线：把上一次实测容量直接当作预测值；它不训练任何参数，是必须跨过的下限。',
              '· 普通最小二乘回归：四个相同输入的线性模型，无正则化、无参数搜索，可写出显式公式。',
              '· 随机森林回归：300 棵树、固定随机种子 25121619，验证段比较 6 组树深与叶大小组合后选参。',
              '· 三者使用完全相同的输入、相同的时间划分、相同的评价指标。',
              '· 每块电池单独建模，指标先逐电池计算再做等权平均；129 个测试目标不代表 129 块独立电池。']},

    {'kind': 'text', 'title': '使用、维护与更换建议', 'size': 16,
     'lines': ['· 保留容量保持作参照基线：它不需要训练，平均误差最低，可作为判断任何模型的下限。',
              '· 逐电池查看预测偏差与最大高估，不能只看平均 RMSE；本次森林后期持续高估，不能据此乐观安排续航。',
              '· 接近研究阈值（首循环容量的 80%）时用新一轮实测复核再决策，不要凭早期线性交点确定更换日期。',
              '· 不要把容量误差换算成飞行分钟数或剩余寿命；数据集缺少任务能耗、电流与内阻信息。',
              '· 需要新电池与更多工况外部验证后才可推广；结论只适用于这四块电池、24℃ 与本次时间划分。']},

    {'kind': 'team', 'title': '成员分工、复现与仓库', 'size': 15,
     'lines': ['· 全仓 37 项自动化测试通过：清洗 7 项、跳变对照 3 项、回归 5 项、森林 4 项、第九步 18 项。',
              '· 运行顺序：01 → 01b → 02 → 06 → 07 → 08 → 09，再执行 python -m unittest discover -s tests。',
              '· 依赖：numpy、pandas、matplotlib、scikit-learn，图表默认使用 Arial Unicode MS。',
              '· 仓库：https://github.com/yxz666-Y/python-midterm-pre',
              '· 结论局限：只有四块电池、单一温度与一次时间划分，尚未做滚动起点重验与多工况外部验证。']},
]

TEAM_TABLE = [
    ['成员', '承担模块', '对应提交'],
    ['杨晓真 25121619', '建模与比较（06/07/08）、报告与过程PPT', '0b358f3 8e4d38e a1a2940'],
    ['待填（本机 xxl-im）', '数据读取与清洗（01/01b）、检查用例与章节', '4efe518 994f659'],
    ['待填（Buln1120）', '第九步评估与不确定性（09）', '90ff04a'],
]


def png_size(path):
    head = Path(path).read_bytes()[:26]
    return struct.unpack('>II', head[16:24])


def style_run(run, size_pt, bold=False):
    run.font.size = Pt(size_pt)
    run.font.name = FONT
    run.font.bold = bold
    properties = run.font._rPr
    element = properties.find(qn('a:ea'))
    if element is None:
        element = parse_xml(f'<a:ea {nsdecls("a")} typeface="{FONT}"/>')
        properties.insert_element_before(element, 'a:cs', 'a:sym', 'a:hlinkClick',
                                         'a:hlinkMouseOver', 'a:rtl', 'a:extLst')
    else:
        element.set('typeface', FONT)


def add_lines(slide, lines, box, size_pt):
    textbox = slide.shapes.add_textbox(*box)
    frame = textbox.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        style_run(run, size_pt)
    return textbox


def add_picture(slide, image, box):
    left, top, max_width, max_height = box
    width, height = png_size(image)
    scale = min(max_width / width, max_height / height)
    return slide.shapes.add_picture(str(image), left, top,
                                    int(width * scale), int(height * scale))


def number_footer(slide, number):
    add_lines(slide, [FOOTER_TEXT], FOOTER_BOX, 11.0)
    add_lines(slide, [str(number)], PAGE_BOX, 11.0)


def build(prs):
    number = 0
    # 封面、汇报内容与数据清洗说明
    for spec in [TEXT_PAGES[0], TEXT_PAGES[1], TEXT_PAGES[2]]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        number += 1
        if spec['kind'] == 'cover':
            add_lines(slide, [spec['title']], (Inches(0.62), Inches(2.30), Inches(12.08), Inches(1.30)), 30.0)
            add_lines(slide, [spec['subtitle']], (Inches(0.62), Inches(3.70), Inches(12.08), Inches(0.70)), 18.0)
            add_lines(slide, spec['lines'], (Inches(0.62), Inches(5.20), Inches(12.08), Inches(1.40)), 16.0)
        else:
            add_lines(slide, [spec['title']], TITLE_BOX, 29.0)
            add_lines(slide, spec['lines'], BODY_BOX, spec['size'])
            number_footer(slide, number)
    # 退化曲线与个体差异、差异量化
    for spec in PAGES[:2]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        number += 1
        add_lines(slide, [spec['title']], TITLE_BOX, 26.0)
        add_picture(slide, spec['figure'], PIC_BOX)
        add_lines(slide, spec['notes'], NOTE_BOX, 15.0)
        add_lines(slide, [spec['caption']], CAPTION_BOX, 13.0)
        number_footer(slide, number)
    # 方法说明 → 结果图 → 建议 → 分工
    for spec in [TEXT_PAGES[3], TEXT_PAGES[4]] + PAGES[2:] + [TEXT_PAGES[5], TEXT_PAGES[6]]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        number += 1
        add_lines(slide, [spec['title']], TITLE_BOX, 26.0)
        if 'figure' in spec:
            add_picture(slide, spec['figure'], PIC_BOX)
            add_lines(slide, spec['notes'], NOTE_BOX, 15.0)
            add_lines(slide, [spec['caption']], CAPTION_BOX, 13.0)
        else:
            box = BODY_BOX
            if spec.get('kind') == 'team':
                box = (Inches(0.68), Inches(3.65), Inches(11.77), Inches(2.90))
                table = slide.shapes.add_table(len(TEAM_TABLE), 3, Inches(0.68), Inches(1.45),
                                               Inches(11.77), Inches(1.90)).table
                for row_index, row in enumerate(TEAM_TABLE):
                    for col_index, cell in enumerate(row):
                        table.cell(row_index, col_index).text = cell
                        for paragraph in table.cell(row_index, col_index).text_frame.paragraphs:
                            for run in paragraph.runs:
                                style_run(run, 14.0, bold=(row_index == 0))
            add_lines(slide, spec['lines'], box, spec['size'])
        number_footer(slide, number)
    return number


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    count = build(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f'已生成 {count} 页：{OUT}')


if __name__ == '__main__':
    main()
