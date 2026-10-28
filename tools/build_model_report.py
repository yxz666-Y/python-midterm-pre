"""生成两轮容量建模与预测专项报告。使用Codex工作区Python依赖运行。"""
from pathlib import Path
import csv
import hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
def read_csv(path):
    with (ROOT/path).open(encoding='utf-8-sig') as f: return list(csv.DictReader(f))
summary=read_csv('results/step01/battery_summary.csv')
common=read_csv('results/step02/common_window_comparison.csv')
metrics=read_csv('results/step03/metrics.csv')
raw_hash=hashlib.sha256((ROOT/'data/battery.csv').read_bytes()).hexdigest()
# 课程模板用于章节结构参考。本阶段报告不填写未经确认的小组身份信息。
doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(19)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
for n in ['Heading 1','Heading 2']:
    doc.styles[n].paragraph_format.space_before=Pt(9)
    doc.styles[n].paragraph_format.space_after=Pt(8)
    doc.styles[n].paragraph_format.keep_with_next=True
# 删除默认模板可能携带的段落边框和主题字体引用。
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if st.type==1 or st.type==2:
        rp=st.element.get_or_add_rPr();rf=rp.rFonts
        if rf is None:
            rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+k),None)
        for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Arial Unicode MS')
# 页码仅用于定位报告内容。
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): return doc.add_heading(t,1)
def page(): doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    if widths:
        for col,w in zip(t.columns,widths): col.width=Cm(w)
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DCE6EC' if ri==0 else ('F3F6F8' if ri%2 else 'FFFFFF'));pr.append(fill)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.1
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(cell.text)<18 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:run.font.size=Pt(9);run.font.bold=ri==0
    rpt=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rpt)
    spacer=p('');spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1;spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)
    return t

def fig(path,caption,width=16):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(ROOT/path),width=Cm(width))
    para.paragraph_format.keep_with_next=True
    p(caption,'Caption')

import json
old_metrics=read_csv('results/step03/metrics.csv')
old_predictions=read_csv('results/step03/predictions.csv')
new_metrics=read_csv('results/step05/metrics.csv')
macro=read_csv('results/step05/macro_common_origins.csv')
old_models=json.loads((ROOT/'results/step03/ridge_models.json').read_text())
new_models=json.loads((ROOT/'results/step05/ridge_models.json').read_text())
new_predictions=read_csv('results/step05/predictions.csv')
batteries=['B0005','B0006','B0007','B0018']
old_names={'persistence':'上一循环基线','linear_cycle':'全训练段线性回归','ridge_history':'历史特征岭回归'}
new_names={'persistence':'容量保持基线','linear_recent20':'最近20次线性趋势','ridge_direct':'直接多步岭回归'}

def metric_old(b,m):return next(r for r in old_metrics if r['battery_id']==b and r['model']==m and r['split']=='test')
def metric_new(b,m,h):return next(r for r in new_metrics if r['battery_id']==b and r['model']==m and int(r['horizon'])==h and r['split']=='test' and r['scope']=='common_origins')
def fmt(x):return f'{float(x):.6f}'
def model_figure(stage,b,caption):fig(f'figures/model_report/step{stage:02}_{b}.png',caption,16.4)

def report_title(t):p(t,'Title')
report_title('两轮电池容量建模与预测报告')
p('杨晓真  25121619\n报告日期  2026年10月28日')
h('主要结果')
p('第一轮单步预测中，上一循环基线在四块电池上的测试MAE和RMSE均最低。第二轮共同起点比较显示，提前1次的四电池平均RMSE以容量保持最低，提前5次和10次以最近20次趋势最低。岭回归在B0006较长距离上有优势，在B0018上则低估明显。以下通过指标、曲线和实际样例说明结果，容量单位均为Ah。')
h('1 两次实验分别解决什么问题')
table(['项目','第一轮单步预测','第二轮多步预测'],[
 ['预测目标','预测下一循环容量','分别预测后1、5、10次容量'],
 ['回归方法','全训练段线性回归、历史特征岭回归','最近20次线性趋势、直接多步岭回归'],
 ['简单参照','沿用上一循环容量','保持预测起点容量'],
 ['主评价样本','129个电池循环目标','每个距离93个共同电池起点']],[3.2,6.6,6.6])
p('第二轮的最近20次线性趋势会在每个预测起点重新拟合，第一轮的全训练段线性回归则在测试前固定系数。二者训练窗口不同，不能因为都叫“线性回归”就当作同一模型。两轮岭回归均为带L2正则化的线性模型，没有使用随机森林或神经网络。')
h('2 数据与退化相关特征')
p('两轮均使用四块电池的636条原始保留版记录，环境温度均为24℃。统计跳变限幅数据只作独立对照，没有替换本报告的训练或测试容量。四个输入是目标循环次数、最近实测容量、最近5次容量均值和最近5次变化斜率。')
p('预测第t次时，第一轮只用第t−1次及以前容量；第二轮在第o次起点预测第o＋h次，只用第o次及以前容量。均值反映近期水平，斜率反映近期变化方向及速度，单位为Ah/循环。恒定温度不作为输入特征。')
page()
h('3 第一轮的划分与模型结果')
p('按每块电池原始记录数的60%和80%向下取整设训练与验证边界。历史特征需要5次观测，因此第一轮从第6次目标开始，形成616个建模样本。参数仅由验证段选择，测试前以训练加验证段重新拟合，模型系数在测试段固定。')
table(['电池','训练目标循环','验证目标循环','测试目标循环'],[
 ['B0005/06/07','6–100','101–134','135–168'],['B0018','6–79','80–105','106–132']],[3.6,4.2,4.2,4.4])
p('普通线性回归仅使用循环次数。岭回归使用全部四个特征，标准化均值与尺度仅来自拟合段，截距不惩罚。alpha候选为0.01、0.1、1、10、100，以验证RMSE最小选取；B0005、B0006、B0007、B0018分别选择1、10、10、1。')
table(['电池','方法','MAE Ah','RMSE Ah'],[[b,old_names[m],fmt(metric_old(b,m)['MAE_Ah']),fmt(metric_old(b,m)['RMSE_Ah'])] for b in batteries for m in old_names],[2.6,6.2,3.8,3.8])
p('MAE表示平均绝对误差，RMSE对较大误差更敏感，二者均越小越好。四块电池上基线均最低，岭回归居中，全训练段线性回归最高。129个目标并非129块独立电池，三种模型的387条预测也不代表387个独立观测。')
page()
h('4 第一轮实测与预测曲线')
p('以下展示全部测试目标，横轴均为目标循环。每次预测可使用此前真实观测，图中整段曲线由多次单步预测组成。')
model_figure(3,'B0005','图1 B0005完整单步测试段 第135至168次')
p('B0005的基线RMSE为0.010001 Ah，岭回归为0.014792 Ah。容量总体下降，但第151次附近存在回升。基线能够跟随局部变化，却通常晚一个循环，固定直线难以跟随这一回升。')
model_figure(3,'B0006','图2 B0006完整单步测试段 第135至168次')
p('B0006的基线RMSE为0.012239 Ah，岭回归为0.015727 Ah，全训练段直线为0.079410 Ah。图中直线在本段持续偏低，说明训练阶段的平均退化速度难以完整描述后期轨迹。这个解释只针对本次时间划分。')
page()
h('4 第一轮实测与预测曲线 续')
model_figure(3,'B0007','图3 B0007完整单步测试段 第135至168次')
p('B0007基线RMSE为0.008210 Ah，在四块电池中最低。局部回升和末尾反弹仍会导致滞后，平均误差低不表示每次容量突变都能提前预测。')
model_figure(3,'B0018','图4 B0018完整单步测试段 第106至132次')
p('B0018基线RMSE为0.027251 Ah，明显高于另外三块。第106次容量由此前1.3567 Ah上升到1.4604 Ah，而三种方法的预测都在1.355 Ah附近，单点低估约0.104 Ah。该突变难以由此前容量直接预知，不能把误差只归结为模型复杂度不够。')
page()
h('5 第二轮多步预测的设计')
p('固定距离h为1、5、10，在每个起点o预测第o＋h次容量。容量保持基线取起点容量；最近20次趋势用该起点及以前最近20次实测容量拟合直线后外推；直接岭回归按每块电池、每个距离独立拟合一个模型，共12个。每次预测内部不使用未来中间循环的真实容量。')
p('模型拟合的目标标签必须在起点时已经可用。以前三块电池为例，验证拟合只用目标不晚于第100次的样本；验证起点从100开始。选参后用目标不晚于134的配对重拟合，测试起点从134开始。不能在第134次起点使用已拟合到第143次容量的模型。')
table(['电池','训练标签截至','测试拟合标签截至','共同测试起点'],[
 ['B0005/06/07','100','134','134–158 各25个'],['B0018','79','105','105–122 共18个']],[3.6,3.6,4.4,4.8])
p('共同测试起点合计93个，使不同距离使用相同的起点集合。各距离对应的目标不同，例如第134次起点的三个目标分别是135、139、144次。同一距离内各方法预测完全相同的目标。')
h('岭回归的参数选择')
table(['电池','提前1次 alpha','提前5次 alpha','提前10次 alpha'],[[b]+[str(new_models[f'{b}_h{h}']['alpha']) for h in [1,5,10]] for b in batteries],[3.4,4.3,4.3,4.4])
p('全部使用相同的5个alpha候选，只按验证RMSE选参，共保存60组搜索。标准化只使用拟合段。测试结果没有用于重新挑选特征、窗口或候选参数。最近20次线性窗口在比较前固定。')
p('第二轮h＝1在“全部合法测试起点”口径下的岭回归预测，与第一轮逐点一致。下一页的共同起点表只保留其中一部分，因此其误差数值与第一轮全测试段表不同，这属于样本范围变化。')
page()
h('6 第二轮逐电池误差与整体表现')
p('以下均为共同测试起点RMSE，单位Ah。前三块各25个起点，B0018为18个。数值越小，表明该组目标上的预测误差越低。')
table(['电池','提前次数','容量保持','最近20次趋势','直接岭回归'],[[b,h]+[fmt(metric_new(b,m,h)['RMSE_Ah']) for m in new_names] for b in batteries for h in [1,5,10]],[2.6,2.5,3.8,3.8,3.7])
p('先对每块电池计算RMSE，再对四块电池等权平均，得到下表。这是四个RMSE的算术平均，并非将全部误差合并后再开方。')
table(['提前次数','容量保持','最近20次趋势','直接岭回归'],[[h]+[fmt(next(r for r in macro if int(r['horizon'])==h and r['model']==m)['RMSE_Ah']) for m in new_names] for h in [1,5,10]],[3.2,4.4,4.4,4.4])
p('平均指标上，容量保持在提前1次时最低，最近20次趋势在5次和10次时最低。个体上，B0006的直接岭回归在5次和10次最低；B0018的岭回归误差较大。统一的平均排名不能替代逐电池检查。')
page()
h('7 第二轮提前10次的预测曲线')
p('以下均为共同起点的提前10次预测，横轴为目标循环。滑动窗口之间相关，误差差异不作为统计显著性证据。')
model_figure(5,'B0005','图5 B0005提前10次 目标为第144至168次')
p('B0005最近20次趋势RMSE为0.017228 Ah，低于容量保持的0.028923 Ah和岭回归的0.039279 Ah。容量保持在整体下降时容易偏高，岭回归在这段则明显偏低。两种偏差方向相反，均需结合实测检验。')
model_figure(5,'B0006','图6 B0006提前10次 目标为第144至168次')
p('B0006直接岭回归RMSE为0.018616 Ah，略低于最近20次趋势的0.021267 Ah，容量保持为0.055006 Ah。趋势信息在该电池的较长距离预测上有帮助，单纯保持当前容量会高估后续容量。')
page()
h('7 第二轮提前10次的预测曲线 续')
model_figure(5,'B0007','图7 B0007提前10次 目标为第144至168次')
p('B0007最近20次趋势RMSE为0.013790 Ah，低于容量保持0.027645 Ah与岭回归0.028233 Ah。末尾容量反弹仍会造成误差，趋势外推不能保证捕捉局部恢复。')
model_figure(5,'B0018','图8 B0018提前10次 目标为第115至132次')
p('B0018直接岭回归RMSE为0.085901 Ah，呈持续低估；最近20次趋势为0.037173 Ah，容量保持为0.045064 Ah。该结果说明当前特征与正则化方案在该段外推受限。需要新增验证再判断成因，不能仅凭曲线断言数据错误或物理退化机制。')
page()
h('8 一个固定起点的实际预测例子')
p('B0006在第134次起点的实测容量为1.3474 Ah。此时分别预测第135、139、144次容量，三种方法都只能使用第134次及以前的观测。下表便于直观看到“预测值”与未来实际值的差别。')
example=[]
for horizon in [1,5,10]:
 rows=[r for r in new_predictions if r['battery_id']=='B0006' and int(r['origin_cycle'])==134 and int(r['horizon'])==horizon and r['split']=='test']
 example.append([horizon,134+horizon,fmt(rows[0]['actual_Ah'])]+[fmt(next(r for r in rows if r['model']==m)['predicted_Ah']) for m in new_names])
table(['提前次数','目标循环','实测 Ah','容量保持','20次趋势','直接岭回归'],example,[2,2.2,3.1,3.1,3,3])
p('这个起点上直接岭回归对第144次预测为1.272258 Ah，实测为1.274100 Ah，误差约−0.001842 Ah；容量保持预测为1.347400 Ah，误差为+0.073300 Ah。这只是一个样例，完整结论仍依据全部共同起点，不能挑选单个成功预测证明普遍有效。')
h('9 如何理解预测结果')
p('第一轮回答“已知上次容量，下一次能预测多准”。第二轮回答“在同一个起点，往后预测更远会怎样”。两轮都完成容量预测，不能直接把误差较小解释为剩余寿命准确，也不能换算真实飞行时长。此前第四步的阈值外推与预警是另一个实验。')
p('容量保持在短期有竞争力，较长距离时可能积累高估。最近20次趋势在本轮较长距离的平均误差更低，但局部恢复仍会造成偏差。岭回归没有在所有电池上稳定胜出，应保留其失败案例和低估风险。现有结果不支持为了提高分数而删除测试中的跳变或更换测试目标。')
p('本次沿用四块、24℃电池和此前已检查过的数据。验证只覆盖时间隔离与当前样本，缺少新电池和多温度外部验证，不构造统计置信区间，也不宣称平均排名已具有普遍性。若进一步调参，应另存版本，并采用新增数据或预先规划的时间验证。')
page()
h('10 数据来源与版本记录')
p('第一轮来源：src/03_features_and_models.py及results/step03。第二轮来源：src/05_multihorizon_models.py及results/step05。报告图只从已保存预测表绘制，没有重训模型。24组第一轮指标和144组第二轮指标已独立重算核对。')
p('第一轮算法版本：f1c88be、1056c97、777282f。第二轮算法与修正版本：b089f67、c624a14、914a154，展示结果cd73031。本报告的数据快照为6ef0b58。输入文件哈希保存在results/model_report/source_verification.json。')
table(['文件或目录','复查用途'],[
 ['results/step03/metrics.csv','第一轮验证及测试MAE、RMSE'],
 ['results/step03/predictions.csv','第一轮每个目标的实测与预测值'],
 ['results/step05/metrics.csv','第二轮按电池、距离、方法和口径的误差'],
 ['results/step05/predictions.csv','第二轮起点、目标、实测与预测值'],
 ['results/step05/macro_common_origins.csv','四块电池等权平均指标'],
 ['两步结果中的ridge_models.json','标准化参数、系数及选定alpha'],
 ['两步结果中的ridge_validation_search.csv','全部候选参数的验证误差'],
 ['figures/model_report/','本报告8张实测与预测曲线图']],[8,8.4])
p('图表生成命令为python tools/plot_model_report.py，仅读取现有结果并独立重算24＋144组指标。生成Word需要tools/requirements-report.txt中的文档依赖，再运行tools/build_model_report.py。')
p('此前36项代码检查已通过，覆盖历史窗口、时间划分、测试隔离和单步跨版本一致性等约束。本次只整理结果和图表，没有修改模型算法、训练参数或测试数据。')
p('仓库 https://github.com/yxz666-Y/python-midterm-pre\n报告由Codex协助编写并核对数值，成员承担按实际理解与核查情况据实说明。')
doc.core_properties.title='两轮电池容量建模与预测报告'
doc.core_properties.author='杨晓真'
output=ROOT/'deliverables/两轮建模与预测报告_杨晓真_25121619.docx'
doc.save(output)
print(output)
