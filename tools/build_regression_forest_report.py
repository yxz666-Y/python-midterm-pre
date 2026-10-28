"""生成普通线性回归与随机森林容量预测报告。使用Codex工作区Python依赖运行。"""
from pathlib import Path
import csv
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
def read_csv(path):
    with (ROOT/path).open(encoding='utf-8-sig') as f: return list(csv.DictReader(f))
# 课程模板用于章节结构参考。本阶段报告不填写未经确认的小组身份信息。
doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(19)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
for n in ['Heading 1','Heading 2']:
    doc.styles[n].paragraph_format.space_before=Pt(9)
    doc.styles[n].paragraph_format.space_after=Pt(8)
    doc.styles[n].paragraph_format.keep_with_next=True
# 删除默认模板可能携带的段落边框和主题字体引用。
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if st.type==1 or st.type==2:
        rp=st.element.get_or_add_rPr();rf=rp.rFonts
        if rf is None:
            rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+k),None)
        for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Arial Unicode MS')
# 页码仅用于定位报告内容。
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): return doc.add_heading(t,1)
def page(): doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    if widths:
        for col,w in zip(t.columns,widths): col.width=Cm(w)
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DCE6EC' if ri==0 else ('F3F6F8' if ri%2 else 'FFFFFF'));pr.append(fill)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.1
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(cell.text)<18 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:run.font.size=Pt(9);run.font.bold=ri==0
    rpt=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rpt)
    spacer=p('');spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1;spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)
    return t

def fig(path,caption,width=16):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(ROOT/path),width=Cm(width))
    para.paragraph_format.keep_with_next=True
    p(caption,'Caption')

import json
batteries=['B0005','B0006','B0007','B0018']
names={'persistence':'容量保持基线','linear_ols':'普通线性回归','random_forest':'随机森林回归'}
test=[r for r in read_csv('results/step08/metrics.csv') if r['split']=='test']
macro=read_csv('results/step08/macro_test.csv')
pooled=read_csv('results/step08/pooled_test.csv')
search=read_csv('results/step07/validation_search.csv')
linear=json.loads((ROOT/'results/step06/linear_models.json').read_text())
forest=json.loads((ROOT/'results/step07/forest_metadata.json').read_text())
example=json.loads((ROOT/'results/step08/worked_example.json').read_text())
ranges=read_csv('results/step08/forest_range_diagnostic.csv')
def fmt(value):return f'{float(value):.6f}'
def metric(b,m):return next(r for r in test if r['battery_id']==b and r['model']==m)

p('普通线性回归与随机森林\n电池容量预测比较报告','Title')
p('杨晓真  学号25121619\n2026年10月28日  题目三 无人设备电池健康评估与剩余寿命预测')
p('本报告比较两种主要模型：实验A使用普通最小二乘线性回归，实验B使用随机森林回归。两者使用相同四个输入、相同时间段和相同下一次容量目标，便于辨认算法差别。结果显示普通回归的平均误差低于本次随机森林，容量保持比较基线的平均误差更低。')
h('1 建模任务与当前方法')
p('我们根据电池已经发生的循环容量变化，预测下一次循环的放电容量。例如第134次充放电结束后，使用截至第134次的真实观测预测第135次容量。第135次实测值在之后获得，只用于评价这一次预测。')
table(['角色','方法','输入信息','预测目标'],[
 ['主要模型A','普通线性回归','四个相同的历史特征','下一次循环容量 Ah'],
 ['主要模型B','随机森林回归','四个相同的历史特征','下一次循环容量 Ah'],
 ['比较基线','容量保持','最近一次实测容量','下一次循环容量 Ah']],[2.8,4.0,5.3,4.3])
p('容量随循环次数的总体下降用于描述退化，近期均值和斜率用于提取容量水平及变化趋势。普通回归把这些信息加权相加，随机森林学习分组规则后平均预测。容量保持不学习参数，仅作为参照，不列为第三种主要模型。')
p('移动到下一个预测目标时，可以加入当时刚测得的历史容量，但两个估计器在整个测试段保持固定。本报告评估单步容量预测，循环次数不能直接换算为天数或飞行分钟，容量误差也不能直接解释为剩余寿命误差。')
page()
h('2 本次两类方法的原理与区别')
p('普通线性回归把四个输入乘以各自系数后相加，再加一个截距，得到预测容量。训练通过最小化预测误差的平方和求出系数。本次是四变量多元线性回归，属于简单、可解释的回归方法，拟合时不对系数增加惩罚项。')
p('随机森林训练300棵回归树。每棵树从拟合段有放回抽样，按特征阈值不断分组，在叶节点取训练容量的平均值。预测时样本在每棵树中走到一个叶节点，再将300棵树的预测平均。它可形成分段的非线性关系。')
table(['比较项','实验A 普通线性回归','实验B 随机森林'],[
 ['拟合关系','截距加四个输入的加权和','多棵决策树的分段预测平均'],
 ['学习内容','各特征对应的回归系数','分裂特征、阈值和叶节点均值'],
 ['关系形状','对四个输入呈线性','可表达非线性和特征交互'],
 ['控制复杂度','本次无正则化或参数搜索','验证段选择树深和叶节点大小'],
 ['容量外推','数值可越过训练容量范围','本次均值型森林受训练容量范围限制'],
 ['解释方式','查看回归系数及误差','查看树路径、平均值及误差']],[3.0,6.7,6.7])
p('严格说，普通回归也是监督学习的一种。本报告按课程要求，用普通线性回归代表简单回归，用树集成代表另一类机器学习方法。两种主要模型都使用同样的输入、时间段和目标，才能对其学习规则及预测表现作直接比较。')
h('同一个样本在随机森林中如何预测')
p(f'以B0006第135次为例，第一棵树先检查上次容量1.3474 Ah是否不超过1.640150 Ah，再检查目标循环135是否大于99，随后继续按阈值分组。该样本最终的第一棵树预测为{fmt(example["first_tree_leaf_value_Ah"])} Ah。')
p(f'300棵树对此样本的预测介于{fmt(example["tree_prediction_min_Ah"])}至{fmt(example["tree_prediction_max_Ah"])} Ah，平均为{fmt(example["tree_prediction_mean_Ah"])} Ah。完整树路径和均值核对保存在worked_example.json中。单棵树输出不是最终森林输出。')
page()
h('3 数据特征与统一评估方法')
p('数据为四块电池636条循环容量记录。两种模型都读取原值保留版清洗结果，不用统计限幅对照替换测试目标。24℃在全部记录中恒定，不能提供样本间差异，因此不作为输入。每块电池单独拟合，前5次只用于构造历史窗口，最终得到616个可建模样本。')
table(['输入特征','预测第t次时的计算方法','B0006预测第135次'],[
 ['目标循环次数','t，预测时已知的循环编号','135'],
 ['上次容量','Q(t−1)','1.347400 Ah'],
 ['最近5次均值','Q(t−5)到Q(t−1)的平均值','1.325340 Ah'],
 ['最近5次斜率','上述5点对循环编号拟合的斜率','0.005130 Ah/循环']],[3.4,8.1,4.9])
p('先把容量序列向后移1位，再计算滚动均值和斜率，避免目标容量进入自己的特征。例如预测第135次时只用第130至134次历史，第135次真实容量1.326000 Ah仅用于事后评分。')
table(['电池','训练目标循环','验证目标循环','测试目标循环'],[
 ['B0005/06/07','6至100 各95条','101至134 各34条','135至168 各34条'],
 ['B0018','6至79 共74条','80至105 共26条','106至132 共27条']],[3.2,4.4,4.4,4.4])
p('两次实验均按时间顺序划分，不随机打乱。训练阶段只拟合训练段。随机森林只使用验证段选择参数，选好后与普通回归一样，在训练加验证段上重拟合。测试期间估计器固定，但每次预测可以使用此前新测得的实际容量。')
p('两种模型各预测相同的129个测试目标，另保留容量保持基线，即预测容量等于上次实测容量。129个循环目标来自四块电池，三种方法的387条预测并不代表387个独立样本。本轮沿用已看过的四块电池，属于追加方法比较。')
h('评价指标')
p('MAE是绝对误差的平均值。RMSE是误差平方平均后开方，对大误差更敏感。二者单位均为Ah，越小越好。平均偏差定义为预测减实测，正数代表高估容量，负数代表低估。')
page()
h('4 实验A 普通线性回归的建模与预测')
p('模型为预测容量＝截距＋循环系数×t＋上次容量系数×Q(t−1)＋均值系数×最近5次均值＋斜率系数×最近5次斜率。代码先按拟合段标准化输入再求最小二乘解，以改善数值尺度，并保存等价的原始单位系数。测试数据不参与标准化。')
p('普通回归不搜索超参数。这里的“简单”指普通回归算法，四个输入共同参与预测，是多元线性回归。普通回归和森林都使用这四个输入，比较时不更换目标或删减困难样本。')
table(['电池','回归MAE Ah','回归RMSE Ah','基线RMSE Ah'],[
 [b,fmt(metric(b,'linear_ols')['MAE_Ah']),fmt(metric(b,'linear_ols')['RMSE_Ah']),fmt(metric(b,'persistence')['RMSE_Ah'])] for b in batteries],[2.8,4.5,4.5,4.6])
p('普通回归在B0006上略优于基线，另外三块上误差更大。它可以跟随近期容量水平和趋势，但并未稳定预测局部突变。不能把加入更多特征理解为误差必然下降。')
h('B0006的最终拟合参数')
model=linear['B0006']
table(['参数','原始单位系数'],[['截距',f'{model["intercept_Ah"]:.9f}']]+[[name,f'{value:.9f}'] for name,value in zip(['循环次数','上次容量','最近5次均值','最近5次斜率'],model['coefficients_raw'])],[8.2,8.2])
p(f'最终模型使用第6至134次的129个拟合目标。代入第3节样本的四个输入，预测B0006第135次容量为{fmt(example["predictions"]["linear_ols"])} Ah，实测为1.326000 Ah。系数完整精度保存在linear_models.json中。')
p('上次容量和近期均值高度相关，因此单个系数不能直接解释为某个因素的独立因果影响。回归系数可以解释计算过程，不能据此推断电池的物理老化机制。实验A版本为a1a2940，已单独提交并推送。')
page()
h('5 实验B 随机森林的选参与外推表现')
p('固定300棵树、有放回抽样、每次分裂允许考察全部4个特征、随机种子25121619。候选树深为2、4、不设上限，候选叶节点最少样本数为3、8，共6种组合。每块电池只用验证RMSE选取组合，共保存24条搜索结果。')
table(['电池','选定最大树深','叶节点最少样本','验证RMSE Ah'],[
 [b,'不限' if forest[b]['selected_parameters']['max_depth'] is None else forest[b]['selected_parameters']['max_depth'],
 forest[b]['selected_parameters']['min_samples_leaf'],
 fmt(min(float(r['RMSE_Ah']) for r in search if r['battery_id']==b))] for b in batteries],[2.8,4.4,4.4,4.8])
p('选参后用训练加验证段重拟合，并把四个模型保存为joblib文件。验证误差来自只用训练段拟合的模型，测试误差来自使用更多已知历史重新拟合的模型，两者不应混为同一拟合状态。实验B版本为8e4d38e。')
h('训练容量范围能解释哪些误差')
table(['电池','拟合容量最小值 Ah','测试容量最小值 Ah','低于拟合最小值的目标数'],[
 [r['battery_id'],fmt(r['fit_capacity_min_Ah']),fmt(r['test_capacity_min_Ah']),f'{r["test_below_fit_min_n"]}/{r["n"]}'] for r in ranges],[2.5,4.5,4.5,4.9])
p('本次森林直接预测绝对容量。叶节点是拟合容量的均值，森林又对叶节点预测求平均，因此输出无法低于全部拟合容量的最小值。前三块电池各有32、33、32个测试容量低于这一最小值，森林在这三块的所有测试目标上均高估容量。')
p('这一边界性质和测试分布的变化，能够解释其难以跟随后期下降。它不表示随机森林在所有电池任务中都无效。以后可另立版本检验预测容量变化量等改进，但本次不看着测试结果追加调参或改写目标。')
p('特征重要性另存为feature_importance.csv，属于训练中不纯度下降的分配。循环、上次容量和均值彼此相关，重要性不代表因果贡献，本报告不据此删除特征。')
page()
h('6 相同测试目标上的误差比较')
table(['电池','方法','MAE Ah','RMSE Ah','平均偏差 Ah'],[
 [b,names[m],fmt(metric(b,m)['MAE_Ah']),fmt(metric(b,m)['RMSE_Ah']),fmt(metric(b,m)['bias_Ah'])] for b in batteries for m in names],[2.1,4.2,3.3,3.3,3.5])
p('前三块各34个测试目标，B0018为27个，方法间目标相同。下表对四块电池的指标等权平均。')
table(['方法','四电池平均MAE Ah','四电池平均RMSE Ah'],[
 [names[m],fmt(next(r for r in macro if r['model']==m)['MAE_Ah']),fmt(next(r for r in macro if r['model']==m)['RMSE_Ah'])] for m in names],[5.4,5.5,5.5])
p('平均RMSE依次为容量保持基线0.014425 Ah、普通回归0.017370 Ah、随机森林0.058766 Ah。平均排名用于描述本次划分，不能取代逐电池判断：B0006普通回归优于基线，B0018随机森林也优于普通回归。')
p('合并129个目标计算的另一种汇总口径另存pooled_test.csv，不与四电池平均指标混用。')
page()
h('7 B0005与B0006的实测预测曲线')
p('横轴为目标循环。下列整段曲线由逐次单步预测组成，移动到下一个目标时会获得此前真实历史，不能理解为第134次一次性预测后续整段。')
fig('figures/step08/B0005_predictions.png','图1 B0005 第135至168次相同测试目标')
p('普通回归能够随最近容量下降，但局部回升仍造成误差。森林曲线停留在约1.37 Ah，难以下降到后段约1.29 Ah的实测水平。容量保持基线通常滞后一个循环，却在该测试段取得更低平均误差。')
fig('figures/step08/B0006_predictions.png','图2 B0006 第135至168次相同测试目标')
p('普通回归在本电池的RMSE为0.011730 Ah，略低于基线的0.012239 Ah。森林预测约1.33 Ah，未跟随后期下降到约1.15至1.19 Ah，其RMSE增至0.100671 Ah。')
page()
h('8 B0007与B0018的实测预测曲线')
fig('figures/step08/B0007_predictions.png','图3 B0007 第135至168次相同测试目标')
p('B0007容量保持基线的RMSE为0.008210 Ah，普通回归为0.011011 Ah。森林接近1.486 Ah的水平线，而实测总体更低，显示相同的外推限制。对末尾容量反弹，单步模型也无法保证提前准确判断。')
fig('figures/step08/B0018_predictions.png','图4 B0018 第106至132次相同测试目标')
p('B0018在第106次出现明显回升，多种方法都低估该点。它的测试容量大多处于先前拟合范围内，森林曲线不再完全平坦，RMSE为0.028542 Ah，低于普通回归0.033624 Ah，但仍略高于基线0.027251 Ah。')
p('同一种森林在不同电池上的表现有差异。本轮平均误差较高的结论，主要结合前三块电池的容量范围变化解释，不能只挑B0018或某个成功点作整体结论。')
page()
h('9 单个预测与全部测试点的对照')
p('B0006第134次结束时，已知上次容量1.347400 Ah。两种新模型都预测第135次容量，真实值在下一次测量后才获得。该例展示两类算法如何对同一个输入给出不同预测。')
table(['方法','预测第135次 Ah','实测 Ah','预测减实测 Ah'],[
 [names[m],fmt(example['predictions'][m]),fmt(example['actual_Ah']),fmt(example['predictions'][m]-example['actual_Ah'])] for m in names],[4.6,4.2,3.4,4.2])
p('此单点上回归和森林都优于基线，但B0006完整测试段的森林误差很大，说明单个成功预测不能代表整体性能。下图展示全部129个目标，虚线为预测等于实测，偏离虚线越远，单点误差越大。')
fig('figures/step08/actual_vs_predicted.png','图5 两类模型的全部测试目标 实测与预测散点')
p('森林在前三块电池上出现近似水平的点带，表示不同的较低实测容量得到了接近的预测值。普通回归更靠近等值线，但B0018仍有较大误差。图中同一电池的相邻点相关，不能把散点数当作独立电池样本数。')
page()
h('10 本轮结论与复现记录')
p('本轮完成了普通最小二乘线性回归和随机森林回归两次新建模。两者使用同样的四个输入、同样的下一循环容量目标和同样的时间划分，差别来自函数形式、学习规则和复杂度控制。普通回归整体优于本次森林，但容量保持基线的四电池平均误差仍最低。')
p('实践上应保留容量保持作参照，逐电池检查预测偏差。对本次直接预测绝对容量的随机森林，前三块电池存在持续高估，不能根据其预测值乐观安排续航。仅凭这四块电池、24℃和单次时间划分，无法给出通用更换标准，也不能换算真实飞行时间。')
p('限制包括测试时间段此前已被查看、只有四块电池、窗口之间相关、缺少新电池与多温度外部验证。普通回归和森林各有失败点，本报告不报告虚构置信区间，也不将容量误差小直接等同于剩余寿命准确。')
table(['记录','内容'],[
 ['a1a2940','实验A代码、完整特征、回归参数、预测与测试'],
 ['8e4d38e','实验B代码、24组验证搜索、四个森林及预测'],
 ['src/08_compare_regression_and_forest.py','重算24组指标、核对相同目标、生成比较图'],
 ['results/step08/verification.json','输入及源码哈希与核对记录']],[6.4,10])
p('依次运行src/01_read_and_clean.py、src/06_linear_regression.py、src/07_random_forest.py、src/08_compare_regression_and_forest.py。特征提取独立保存在src/comparison_common.py中。使用requirements.txt的分析依赖。安装tools/requirements-report.txt后运行tools/build_regression_forest_report.py生成本文。')
p('当前12项代码检查通过，覆盖统计跳变对照、历史特征不读取当前或未来容量、时间划分、回归参数重建、森林选参及保存模型复现等约束。24组误差指标已从逐点预测独立重算，两种模型和比较基线的目标与实测值一致，实测值与原始CSV逐点一致。')
p('方法参考为scikit-learn官方文档：LinearRegression说明普通最小二乘；RandomForestRegressor说明回归树集成和平均预测。实现使用scikit-learn 1.7.2，森林参数配置另存experiment.json。')
p('https://scikit-learn.org/1.7/modules/generated/sklearn.linear_model.LinearRegression.html\nhttps://scikit-learn.org/1.7/modules/generated/sklearn.ensemble.RandomForestRegressor.html')
p('仓库 https://github.com/yxz666-Y/python-midterm-pre\n报告由Codex协助编写和核对，成员承担按实际理解与核查情况据实说明。')
target=ROOT/'deliverables/普通线性回归与随机森林预测报告_杨晓真_25121619.docx'
doc.save(target)
print(target)

