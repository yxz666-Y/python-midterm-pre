"""生成阶段性数据分析报告。使用Codex工作区Python依赖运行。"""
from pathlib import Path
import csv
import hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
def read_csv(path):
    with (ROOT/path).open(encoding='utf-8-sig') as f: return list(csv.DictReader(f))
summary=read_csv('results/step01/battery_summary.csv')
common=read_csv('results/step02/common_window_comparison.csv')
metrics=read_csv('results/step03/metrics.csv')
raw_hash=hashlib.sha256((ROOT/'data/battery.csv').read_bytes()).hexdigest()
# 课程模板用于章节结构参考。本阶段报告不填写未经确认的小组身份信息。
doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(19)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
for n in ['Heading 1','Heading 2']:
    doc.styles[n].paragraph_format.space_before=Pt(9)
    doc.styles[n].paragraph_format.space_after=Pt(8)
    doc.styles[n].paragraph_format.keep_with_next=True
# 删除默认模板可能携带的段落边框和主题字体引用。
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if st.type==1 or st.type==2:
        rp=st.element.get_or_add_rPr();rf=rp.rFonts
        if rf is None:
            rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+k),None)
        for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Arial Unicode MS')
# 页码仅用于定位报告内容。
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): return doc.add_heading(t,1)
def page(): doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DCE6EC' if ri==0 else ('F3F6F8' if ri%2 else 'FFFFFF'));pr.append(fill)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.1
                for run in para.runs:run.font.size=Pt(9);run.font.bold=ri==0
    rpt=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rpt)
    spacer=p('');spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1;spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)
    return t

def fig(path,caption,width=16):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(ROOT/path),width=Cm(width))
    para.paragraph_format.keep_with_next=True
    p(caption,'Caption')

p('无人设备电池容量退化数据分析报告','Title')
p('杨晓真  25121619\n分析日期  2026年10月28日')
h('摘要')
p('基于课程提供的636条放电循环记录，本报告分析四块电池的容量退化、统计跳变、单步容量预测及阈值外推风险。数据无缺失和重复记录，但容量并非单调下降。统一前132次循环后，B0006容量下降35.36%，B0007下降21.90%，相差13.46个百分点。变化量检测标记58条待核查观测，其中30条极端跳变另做限幅对照。当前时间划分下，上一循环容量基线的测试误差最低。早期寿命外推对观察阶段敏感，单步容量误差较小也不能直接证明提前预警有效。')
p('关键词  电池容量  退化曲线  统计跳变  时间划分  容量预测')
h('1 引言')
p('电池容量衰减会影响无人设备的可用能量。本项目先检查数据质量并比较个体差异，再以历史容量预测后续容量，核查模型误差与阈值事件。本报告使用放电容量作为观测指标。数据不包含任务能耗、内阻或电流电压时序，因此不换算飞行分钟数，也不推断事故概率。')
h('2 数据介绍')
p('数据文件为data/battery.csv。课程说明称该数据整理自NASA PCoE公开电池老化数据集。本次实际分析对象为所提供CSV，未对照NASA原始实验档案逐条核验。字段为battery_id、cycle、capacity_Ah和ambient_temp_C，环境温度均为24℃。')
table(['电池','记录数','首循环容量 Ah','末循环容量 Ah'],[[r['battery_id'],r['records'],f"{float(r['first_capacity_Ah']):.4f}",f"{float(r['last_capacity_Ah']):.4f}"] for r in summary],[3,2,5.7,5.7])
p('前三块电池各168次，B0018仅132次。总体容量范围为1.1538至2.0353 Ah。观测终止不能等同于电池失效。')
page()
h('3 数据清洗与预处理')
p('使用pandas读取UTF-8 CSV，检查列名、行列数和类型。编号去除首尾空白，循环、容量与温度转换为数值。随后检查缺失、完全重复、同一电池同一循环的冲突、非有限数值、非正容量和循环连续性。按电池与循环排序，原始文件独立保留。')
table(['检查项目','结果','处理'],[['缺失值及数值转换失败','0','无需填充'],['完全重复及主键冲突','0','无记录删除'],['非正容量或非有限数值','0','无需修正'],['循环缺口','0','各电池从1连续编号']],[6,2,8.4])
p('基础检查后仍保留636条观测。“没有明显格式错误”不能写成“没有任何异常值”。为避免把正常后期低容量当作异常，补充检查相邻容量变化ΔQ，而非直接对容量高低作全局筛选。')
p('每块电池用前60%记录的ΔQ估计Q1、Q3和IQR。超出[Q1−1.5IQR，Q3+1.5IQR]标记待核查，超出3IQR范围标记极端。阈值不使用验证或测试数据。训练段内部标记是事后诊断，不能称为在线预警。')
table(['电池','待核查条数','其中极端条数'],[['B0005',9,6],['B0006',26,10],['B0007',9,5],['B0018',14,9],['合计',58,30]],[5.4,5.5,5.5])
p('限幅对照保留全部行，只把30个极端变化量裁剪到3IQR边界，再加回上一循环原始容量。未命中的606条不变，正式清洗数据也保留原值。该处理可能压低真实容量恢复峰值，不能证明修正值更接近真值。')
table(['电池与循环','原容量 Ah','ΔQ Ah','限幅对照 Ah'],[['B0006 第90次','1.593600','+0.151900','1.467350'],['B0018 第46次','1.726700','+0.131200','1.615675']],[4.4,4,4,4])
p('四块电池处理前后首末容量均未改变，全段容量线性斜率仍为负。现有模型使用保留原值的清洗数据，限幅版留作敏感性比较，尚未验证其预测收益。')
page()
h('4 退化曲线与个体差异')
fig('figures/step02/01_capacity_all_cycles.png','图1 全部观测容量曲线 B0018在第132次结束',16)
p('四块电池总体退化，但具有局部回升。B0006首循环容量最高，后期却低于另外两块168次电池。不能只按初始容量判断耐久性，也不能因短期回升就认定数据错误。')
p('为统一观察长度，取共同的第1至132次循环，并将各自首循环容量归一为100%。此“相对首循环容量”与基于厂商额定容量定义的SOH口径不同。下降比例为1−Q132/Q1。')
table(['电池','第132次容量 Ah','容量下降比例','回升间隔数'],[[r['battery_id'],f"{float(r['end_capacity_Ah']):.4f}",f"{float(r['loss_pct']):.2f}%",r['rebound_count']] for r in common],[3.4,4.6,4.4,4])
p('相同132次窗口下，B0006下降最多，B0007最少。B0005与B0018下降比例接近，但这不代表后续寿命相同。回升次数均基于131个相邻间隔统计，不能将次数直接当作健康排名，也不能由现有字段确认差异的物理原因。')
page()
h('5 特征提取与已完成的单步建模')
p('预测第t次容量时，只使用此前已经获得的观测。输入为目标循环次数、上一循环容量、此前5次容量均值和5点线性斜率。先shift(1)再rolling(5)，前5条缺少完整历史，636条观测形成616个建模样本。')
p('逐电池按原始记录数60%和80%向下取整划分。B0005、B0006、B0007的训练目标为6至100次，验证为101至134次，测试为135至168次。B0018对应6至79、80至105、106至132次。参数只用验证段选择，测试前在训练加验证段重拟合，标准化参数仅来自拟合段。')
p('方法包括上一循环基线、仅以循环次数为输入的线性回归，以及历史特征岭回归。岭回归alpha候选固定为0.01、0.1、1、10、100，用验证RMSE选择。测试时逐次获得新的历史实测值，模型系数固定。')
table(['电池','基线 RMSE','线性回归 RMSE','岭回归 RMSE'],[[b]+[f"{float(next(r for r in metrics if r['battery_id']==b and r['split']=='test' and r['model']==m)['RMSE_Ah']):.6f}" for m in ['persistence','linear_cycle','ridge_history']] for b in ['B0005','B0006','B0007','B0018']],[3.4,4.2,4.4,4.4])
p('表中单位均为Ah。三种方法评价同一批129个电池循环目标。基线在四块电池上的测试MAE和RMSE均最低，岭回归优于循环线性回归。复杂方法未胜出，应保留这一结果，避免根据测试表现反复修改参数。')
fig('figures/step03/model_comparison.png','图2 当前单步预测的测试误差比较',15)
page()
h('6 预测风险与寿命外推分析')
p('误差定义为预测减实测。正值表示容量高估，负值表示低估。基线虽有较低RMSE，容量高估比例仍为64.7%至82.4%。测试绝对误差90%分位数是描述统计，不能作为未来90%覆盖率的保证。')
p('第四步以首循环容量的80%作为研究阈值，首次实测容量小于或等于该值定义为首次跨越。该口径没有持续跨越确认，不能等同于不可逆失效或设备通用更换标准。')
table(['电池','首次跨越循环','第三步测试起点'],[['B0005',101,135],['B0006',61,135],['B0007',124,135],['B0018',75,106]],[5.4,5.5,5.5])
p('四块电池的首次跨越均早于测试段，因此第三步结果不能验证首次跨越前的预警能力。低容量状态分类指标即便为1，也受到测试样本全部低于阈值的限制。')
p('组员新增的探索性滚动实验从第30次起，用最近20次容量线性趋势预测未来5次的首次跨越。B0006提前1次检出，B0007与B0018提前5次检出，B0005漏报；后两块分别出现5和3个误报起点。每种方法241个起点仍只对应4块电池的4次首次事件，重叠窗口不能视为独立失效样本。')
table(['电池','前30%历史预测到达循环','前50%历史预测到达循环','实测首次跨越'],[['B0005','228.61','115.09','101'],['B0006','68.57','已跨越','61'],['B0007','216.71','121.01','124'],['B0018','63.00','83.14','75']],[3.2,5,5,3.2])
p('固定起点外推对历史长度敏感。B0005早期估计比实测晚约127.61次，B0007也存在明显高估。不同窗口的预测差异属于敏感性分析，不能称为置信区间。四块电池、24℃及同一数据上的探索限制了推广范围。')
page()
h('7 总结与后续建模方案')
p('数据支持电池间容量退化差异，也显示整体下降与局部恢复并存。建议持续分电池记录容量和预测偏差，在容量突变或接近研究阈值时复测，结合任务能耗、设备规范和其他健康指标决定维护。现有数据无法确定统一更换循环数或具体安全余量。')
p('下一轮在保留原值的数据上比较提前1、5、10次的容量保持、最近20次线性趋势与直接多步岭回归。拟合目标不得晚于预测起点，参数仅用验证段选择。')
p('每个起点内不读取未来实测容量，下一起点才更新历史。实验沿用此前已分析的四块电池，属于探索性扩展，需新增数据检验。')
h('8 已核实的协作记录')
table(['身份记录','已核实的仓库工作','证据'],[['杨晓真 25121619','读取清洗、退化分析、单步模型、统计跳变对照及过程材料','f1c88be、1056c97、777282f、6aa6676等'],['GitHub账号 Buln1120','第四步误差、寿命敏感性及滚动预警评估','8b282e4、ddec284'],['其他成员','本报告未确认其具体提交或分工','需由本人据实补充']],[3.9,7,5.5])
p('表格依据Git贡献记录整理，账号与成员真实姓名的对应关系需小组确认。杨晓真的本轮工作由Codex协助编写代码、运行检查和整理文档，个人承担应结合本人实际核查与理解如实说明。')
h('9 附录与复现')
p('当前数据与结果基于仓库60d0309，原始CSV的SHA256为：')
p(raw_hash)
p('运行次序：01_read_and_clean.py；02_visualize_degradation.py；01b_check_capacity_jumps.py；03_features_and_models.py；04_evaluation_and_uncertainty.py。检查命令为python -m unittest discover -s tests -v，本轮27项既有检查通过。')
p('结果目录对应results/step01、step01b、step02、step03、step04。每个算法或参数方案调整使用独立Git提交并推送，保留原始数据、输入口径、搜索记录和结果。')
p('仓库 https://github.com/yxz666-Y/python-midterm-pre\n课程结构参考 桌面python目录中的项目报告模版.docx\n题目与数据来源 课程题目三说明及所提供CSV')
doc.core_properties.title='无人设备电池容量退化数据分析报告'
doc.core_properties.author='杨晓真'
output=ROOT/'deliverables/数据分析报告_杨晓真_25121619.docx'
doc.save(output)
print(output)
